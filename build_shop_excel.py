#!/usr/bin/env python3
"""쇼핑몰(에이블리 파트너스) 사입·지출 관리 엑셀 생성 스크립트.

실행: python3 build_shop_excel.py
결과: OUTPUT_PATH 에 지정된 xlsx 파일 생성 (거래처/지출/지인판매/에이블리정산/리뷰협찬 5개 시트)
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# 설정값 (세율/행 수 등 — 구조 변경 시 여기만 수정하면 됨)
# ---------------------------------------------------------------------------

OUTPUT_PATH = Path(__file__).parent / "쇼핑몰_사입_지출_관리.xlsx"

FONT_NAME = "맑은 고딕"
HEADER_BG = "2F2F2F"
BAND_BG = "F5F5F5"
MONEY_FORMAT = '#,##0"원"'
DATE_FORMAT = "yyyy-mm-dd"
MONTH_FORMAT = "yyyy-mm"

TAX_RATE_지출_세금계산서 = 0.10
TAX_RATE_지인판매 = 0.20
수수료율_에이블리 = 0.10

ROW_COUNTS = {
    "거래처 관리": 30,
    "지출 관리": 40,
    "지인 판매 관리": 30,
    "에이블리 정산 관리": 24,
    "리뷰 협찬 관리": 30,
}

# ---------------------------------------------------------------------------
# 스타일 유틸
# ---------------------------------------------------------------------------

thin_side = Side(style="thin", color="D0D0D0")
THIN_BORDER = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill("solid", fgColor=HEADER_BG)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

BODY_FONT = Font(name=FONT_NAME, size=10)
BAND_FILL = PatternFill("solid", fgColor=BAND_BG)
BODY_ALIGN = Alignment(horizontal="center", vertical="center")


def write_title_block(ws: Worksheet, title: str, subtitle: str, memo: str | None, last_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    c1 = ws.cell(row=1, column=1, value=title)
    c1.font = Font(name=FONT_NAME, bold=True, size=14)
    c1.alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    c2 = ws.cell(row=2, column=1, value=subtitle)
    c2.font = Font(name=FONT_NAME, italic=True, size=9, color="808080")
    c2.alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    c3 = ws.cell(row=3, column=1, value=memo or "")
    c3.font = Font(name=FONT_NAME, size=9, color="C00000")
    c3.alignment = Alignment(horizontal="left", vertical="center")


def write_header(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


def style_body_rows(ws: Worksheet, header_row: int, row_count: int, last_col: int) -> None:
    for i in range(row_count):
        row = header_row + 1 + i
        banded = i % 2 == 1
        for col in range(1, last_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = BODY_FONT
            cell.alignment = BODY_ALIGN
            cell.border = THIN_BORDER
            if banded:
                cell.fill = BAND_FILL


def add_no_column(ws: Worksheet, header_row: int, row_count: int, anchor_col_letter: str = "B") -> None:
    for i in range(row_count):
        row = header_row + 1 + i
        ws.cell(
            row=row,
            column=1,
            value=f'=IF({anchor_col_letter}{row}="","",COUNTA($B${header_row + 1}:{anchor_col_letter}{row}))',
        )


def add_dropdown(ws: Worksheet, col_letter: str, header_row: int, row_count: int, options: list[str]) -> None:
    dv = DataValidation(type="list", formula1='"{}"'.format(",".join(options)), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{header_row + 1}:{col_letter}{header_row + row_count}")


def set_column_formats(ws: Worksheet, header_row: int, row_count: int, col_letter: str, number_format: str) -> None:
    start = header_row + 1
    end = header_row + row_count
    for row in range(start, end + 1):
        ws[f"{col_letter}{row}"].number_format = number_format


def add_sum_row(ws: Worksheet, sum_row: int, label_col: int, money_cols: list[str], header_row: int, row_count: int) -> None:
    ws.cell(row=sum_row, column=label_col, value="합계")
    start = header_row + 1
    end = header_row + row_count
    for col_letter in money_cols:
        cell = ws[f"{col_letter}{sum_row}"]
        cell.value = f"=SUM({col_letter}{start}:{col_letter}{end})"
        cell.number_format = MONEY_FORMAT
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=sum_row, column=col)
        cell.font = Font(name=FONT_NAME, bold=True, size=10)
        cell.border = THIN_BORDER
        cell.alignment = BODY_ALIGN


def set_widths(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------------------
# 시트 1. 거래처 관리
# ---------------------------------------------------------------------------

def build_거래처관리(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "거래처 관리"
    headers = [
        "No", "거래처명", "위치정보(주소)", "담당자/연락처", "개인평가",
        "평가 메모", "계좌번호(은행)", "세금계산서 발행 여부", "세금계산서 발행일",
        "샘플 가능 여부", "비고",
    ]
    last_col = len(headers)
    header_row = 4
    row_count = ROW_COUNTS["거래처 관리"]

    write_title_block(
        ws, "거래처 관리", "사입처 정보 관리 대장",
        "", last_col,
    )
    write_header(ws, header_row, headers)
    style_body_rows(ws, header_row, row_count, last_col)
    add_no_column(ws, header_row, row_count)

    add_dropdown(ws, "E", header_row, row_count, ["★1", "★2", "★3", "★4", "★5"])
    add_dropdown(ws, "H", header_row, row_count, ["필요", "불필요", "요청완료", "발행완료"])
    add_dropdown(ws, "J", header_row, row_count, ["가능", "불가능", "조건부가능"])
    set_column_formats(ws, header_row, row_count, "I", DATE_FORMAT)

    set_widths(ws, [6, 16, 24, 18, 10, 22, 18, 16, 14, 14, 24])
    ws.freeze_panes = f"A{header_row + 1}"


# ---------------------------------------------------------------------------
# 시트 2. 지출 관리
# ---------------------------------------------------------------------------

def build_지출관리(wb: Workbook) -> None:
    ws = wb.create_sheet("지출 관리")
    headers = [
        "No", "날짜", "지출항목", "거래처명", "지출금액",
        "세금계산서 발행 필요여부", "세금(10%) 입금액", "총 지급액",
        "세금계산서 발행여부", "결제수단", "비고",
    ]
    last_col = len(headers)
    header_row = 4
    row_count = ROW_COUNTS["지출 관리"]
    sum_row = header_row + 1 + row_count

    write_title_block(
        ws, "지출 관리", "샘플비·광고비·사입 교통비 등 지출 및 세금계산서 발행 관리",
        "※ 세금계산서를 요청하려면 거래처에 지출금액의 10%에 해당하는 세금을 별도 입금해야 발행됩니다.",
        last_col,
    )
    write_header(ws, header_row, headers)
    style_body_rows(ws, header_row, row_count, last_col)
    add_no_column(ws, header_row, row_count)

    add_dropdown(ws, "C", header_row, row_count, ["샘플비", "광고비", "사입 교통비", "포장·부자재", "기타"])
    add_dropdown(ws, "F", header_row, row_count, ["필요", "불필요"])
    add_dropdown(ws, "I", header_row, row_count, ["미발행", "요청완료", "발행완료", "해당없음"])
    set_column_formats(ws, header_row, row_count, "B", DATE_FORMAT)

    start, end = header_row + 1, header_row + row_count
    for row in range(start, end + 1):
        ws[f"G{row}"] = f'=IF(F{row}="필요", E{row}*{TAX_RATE_지출_세금계산서}, 0)'
        ws[f"H{row}"] = f"=E{row}+G{row}"
        for col_letter in ("E", "G", "H"):
            ws[f"{col_letter}{row}"].number_format = MONEY_FORMAT

    add_sum_row(ws, sum_row, 4, ["E", "G", "H"], header_row, row_count)

    set_widths(ws, [6, 12, 14, 16, 12, 16, 14, 12, 14, 12, 20])
    ws.freeze_panes = f"A{header_row + 1}"


# ---------------------------------------------------------------------------
# 시트 3. 지인 판매 관리
# ---------------------------------------------------------------------------

def build_지인판매관리(wb: Workbook) -> None:
    ws = wb.create_sheet("지인 판매 관리")
    headers = [
        "No", "날짜", "지인명", "상품명", "도매가",
        "세금(20%)", "판매가(도매가+세금)", "결제여부", "비고",
    ]
    last_col = len(headers)
    header_row = 4
    row_count = ROW_COUNTS["지인 판매 관리"]
    sum_row = header_row + 1 + row_count

    write_title_block(
        ws, "지인 판매 관리", "지인에게 도매가+세금 포함 금액으로 판매하는 거래 관리",
        "※ 판매가 = 도매가 + 도매가의 20% (에이블리를 거치지 않는 거래로, 추후 종합소득세 납부 대비 안전 버퍼 반영)",
        last_col,
    )
    write_header(ws, header_row, headers)
    style_body_rows(ws, header_row, row_count, last_col)
    add_no_column(ws, header_row, row_count)

    add_dropdown(ws, "H", header_row, row_count, ["입금완료", "미입금"])
    set_column_formats(ws, header_row, row_count, "B", DATE_FORMAT)

    start, end = header_row + 1, header_row + row_count
    for row in range(start, end + 1):
        ws[f"F{row}"] = f"=E{row}*{TAX_RATE_지인판매}"
        ws[f"G{row}"] = f"=E{row}+F{row}"
        for col_letter in ("E", "F", "G"):
            ws[f"{col_letter}{row}"].number_format = MONEY_FORMAT

    add_sum_row(ws, sum_row, 3, ["E", "F", "G"], header_row, row_count)

    set_widths(ws, [6, 12, 14, 18, 12, 12, 18, 12, 20])
    ws.freeze_panes = f"A{header_row + 1}"


# ---------------------------------------------------------------------------
# 시트 4. 에이블리 정산 관리
# ---------------------------------------------------------------------------

def build_에이블리정산관리(wb: Workbook) -> None:
    ws = wb.create_sheet("에이블리 정산 관리")
    headers = [
        "No", "정산월", "판매금액(매출)", "에이블리 수수료(10%)", "정산 예정금액(90%)",
        "원천징수세(필요시 입력)", "실수령액", "정산 여부", "비고",
    ]
    last_col = len(headers)
    header_row = 4
    row_count = ROW_COUNTS["에이블리 정산 관리"]
    sum_row = header_row + 1 + row_count

    write_title_block(
        ws, "에이블리 정산 관리", "에이블리 파트너스가 판매금의 10%를 수수료로 수취하는 구조 정산 관리",
        "※ 원천징수 3.3%는 선납 개념이며, 실제 세액은 다음 해 5월 종합소득세 신고 시 확정됩니다 (참고용).",
        last_col,
    )
    write_header(ws, header_row, headers)
    style_body_rows(ws, header_row, row_count, last_col)
    add_no_column(ws, header_row, row_count)

    add_dropdown(ws, "H", header_row, row_count, ["정산대기", "정산완료"])
    set_column_formats(ws, header_row, row_count, "B", MONTH_FORMAT)

    start, end = header_row + 1, header_row + row_count
    for row in range(start, end + 1):
        ws[f"D{row}"] = f"=C{row}*{수수료율_에이블리}"
        ws[f"E{row}"] = f"=C{row}-D{row}"
        ws[f"G{row}"] = f"=E{row}-F{row}"
        for col_letter in ("C", "D", "E", "F", "G"):
            ws[f"{col_letter}{row}"].number_format = MONEY_FORMAT

    add_sum_row(ws, sum_row, 2, ["C", "D", "E", "F", "G"], header_row, row_count)

    set_widths(ws, [6, 12, 16, 16, 16, 18, 14, 12, 20])
    ws.freeze_panes = f"A{header_row + 1}"


# ---------------------------------------------------------------------------
# 시트 5. 리뷰 협찬 관리
# ---------------------------------------------------------------------------

def build_리뷰협찬관리(wb: Workbook) -> None:
    ws = wb.create_sheet("리뷰 협찬 관리")
    headers = [
        "No", "구분", "이름/계정", "연락처", "플랫폼", "협찬 상품",
        "리뷰어 결제금액", "페이백 금액(결제금액 전액)", "리뷰 링크", "리뷰 게시일",
        "리뷰 작성 여부", "페이백 여부", "삭제금지 종료일", "삭제금지 상태", "비고",
    ]
    last_col = len(headers)
    header_row = 4
    row_count = ROW_COUNTS["리뷰 협찬 관리"]
    sum_row = header_row + 1 + row_count

    write_title_block(
        ws, "리뷰 협찬 관리", "지인/인플루언서 리뷰 협찬, 페이백, 최소 3개월 삭제금지 계약 관리",
        "※ 리뷰어 결제 → 리뷰 게시 확인 → 결제금액 전액 페이백. 리뷰 게시일 기준 최소 3개월간 삭제 금지.",
        last_col,
    )
    write_header(ws, header_row, headers)
    style_body_rows(ws, header_row, row_count, last_col)
    add_no_column(ws, header_row, row_count)

    add_dropdown(ws, "B", header_row, row_count, ["지인", "인플루언서"])
    add_dropdown(ws, "E", header_row, row_count, ["인스타그램", "블로그", "유튜브", "기타"])
    add_dropdown(ws, "K", header_row, row_count, ["완료", "미완료"])
    add_dropdown(ws, "L", header_row, row_count, ["완료", "미완료"])
    set_column_formats(ws, header_row, row_count, "J", DATE_FORMAT)
    set_column_formats(ws, header_row, row_count, "M", DATE_FORMAT)

    start, end = header_row + 1, header_row + row_count
    for row in range(start, end + 1):
        ws[f"H{row}"] = f"=G{row}"
        ws[f"M{row}"] = f'=IF(J{row}="","",EDATE(J{row},3))'
        ws[f"N{row}"] = f'=IF(J{row}="","",IF(TODAY()<M{row},"삭제금지 기간","기간 만료(삭제가능)"))'
        for col_letter in ("G", "H"):
            ws[f"{col_letter}{row}"].number_format = MONEY_FORMAT

    add_sum_row(ws, sum_row, 3, ["G", "H"], header_row, row_count)

    set_widths(ws, [6, 10, 14, 14, 12, 16, 14, 18, 16, 12, 12, 12, 14, 20, 20])
    ws.freeze_panes = f"A{header_row + 1}"


# ---------------------------------------------------------------------------
# 실행
# ---------------------------------------------------------------------------

def main() -> None:
    wb = Workbook()
    build_거래처관리(wb)
    build_지출관리(wb)
    build_지인판매관리(wb)
    build_에이블리정산관리(wb)
    build_리뷰협찬관리(wb)

    wb.save(OUTPUT_PATH)
    print(f"생성 완료: {OUTPUT_PATH} ({date.today().isoformat()})")


if __name__ == "__main__":
    main()
