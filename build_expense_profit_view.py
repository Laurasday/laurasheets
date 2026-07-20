#!/usr/bin/env python3
"""지출 관리 + 월별 순수익금 뷰 시트 생성 스크립트.

- 지출 관리: build_shop_excel.py의 템플릿을 그대로 재사용 (빈 템플릿, 40행)
- 월별 순수익금: 같은 파일의 지출 관리 탭 + 마스터 파일(거래처 관리 등)의
  에이블리 파트너스 정산 관리 탭을 IMPORTRANGE로 실시간 연동해 월별 순수익 계산

실행: python3 build_expense_profit_view.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import build_shop_excel as base

OUTPUT_PATH = Path(__file__).parent / "지출관리_월별순수익금.xlsx"

MASTER_SHEET_URL = "https://docs.google.com/spreadsheets/d/1g8rE93aEcWroOBt0uoqezmtkRysxi-Wx3-qcBD4c0zE/edit"
MASTER_SETTLEMENT_TAB = "에이블리 파트너스 정산 관리"
# 마스터 파일의 정산 데이터 범위: 헤더 5행, 데이터 6~29행 (No,정산월,판매금액,수수료,정산예정금액,원천징수세,실수령액,정산여부,비고)
MASTER_SETTLEMENT_DATA_RANGE = f"{MASTER_SETTLEMENT_TAB}!B6:G29"

MASTER_VENDOR_TAB = "거래처 관리 대장"
# 마스터 파일의 거래처 데이터 범위: 헤더 4행, 데이터 5~34행 (No~비고, 12열)
MASTER_VENDOR_DATA_RANGE = f"{MASTER_VENDOR_TAB}!A5:L34"
거래처_HEADERS = [
    "No", "거래처명", "위치정보(주소)", "담당자/연락처", "메신저",
    "개인평가 (★1~5)", "평가 메모", "계좌번호(은행)",
    "세금계산서 발행 여부", "세금계산서 발행일", "샘플 가능 여부", "비고",
]

# 지출 관리 탭 내부 구조 (build_shop_excel.build_지출관리 기준)
EXPENSE_HEADER_ROW = 4
EXPENSE_ROW_COUNT = base.ROW_COUNTS["지출 관리"]
EXPENSE_DATE_COL = "B"
EXPENSE_TOTAL_PAID_COL = "H"
EXPENSE_DATA_START = EXPENSE_HEADER_ROW + 1
EXPENSE_DATA_END = EXPENSE_HEADER_ROW + EXPENSE_ROW_COUNT

MONTHS_IN_YEAR = 2026


def build_월별순수익금(wb: Workbook) -> None:
    ws = wb.create_sheet("월별 순수익금")

    last_col = 4
    base.write_title_block(
        ws, "월별 순수익금", "지출 관리 + 에이블리 정산 실수령액(원본 파일 실시간 연동)을 월별로 집계",
        "※ 수익(정산금)은 원본 거래처 관리 파일의 '에이블리 파트너스 정산 관리' 탭을 IMPORTRANGE로 실시간으로 가져옵니다. 최초 1회 '액세스 허용' 클릭이 필요합니다.",
        last_col,
    )

    header_row = 4
    headers = ["월", "지출 합계", "수익(정산금) 합계", "순수익금"]
    base.write_header(ws, header_row, headers)

    row_count = 12
    base.style_body_rows(ws, header_row, row_count, last_col)

    # 원본 파일의 에이블리 정산 데이터를 헬퍼 영역(F열)에 한 번만 IMPORTRANGE로 가져옴
    helper_row = header_row
    ws.cell(row=helper_row, column=6, value="(참고용 · 원본 연동)")
    ws.cell(row=helper_row, column=6).font = Font(name=base.FONT_NAME, italic=True, size=8, color="808080")
    ws.cell(
        row=helper_row + 1,
        column=6,
        value=f'=IMPORTRANGE("{MASTER_SHEET_URL}","{MASTER_SETTLEMENT_DATA_RANGE}")',
    )

    helper_month_col = "F"   # 정산월 (마스터 B열)
    helper_income_col = "K"  # 실수령액 (마스터 G열 = 헬퍼 F열부터 6번째 → K열)
    helper_start = helper_row + 1
    helper_end = helper_row + 1 + 23  # 24행(6~29행) 스필

    start, end = header_row + 1, header_row + row_count
    for i, row in enumerate(range(start, end + 1)):
        month_index = i + 1
        ws.cell(row=row, column=1, value=f"=DATE({MONTHS_IN_YEAR},{month_index},1)")
        ws.cell(row=row, column=1).number_format = base.MONTH_FORMAT

        ws.cell(
            row=row,
            column=2,
            value=(
                f'=SUMPRODUCT((TEXT(\'지출 관리\'!${EXPENSE_DATE_COL}${EXPENSE_DATA_START}:'
                f'${EXPENSE_DATE_COL}${EXPENSE_DATA_END},"yyyy-mm")=TEXT($A{row},"yyyy-mm"))'
                f'*\'지출 관리\'!${EXPENSE_TOTAL_PAID_COL}${EXPENSE_DATA_START}:'
                f'${EXPENSE_TOTAL_PAID_COL}${EXPENSE_DATA_END})'
            ),
        )
        ws.cell(
            row=row,
            column=3,
            value=(
                f'=SUMPRODUCT((TEXT(${helper_month_col}${helper_start}:${helper_month_col}${helper_end},"yyyy-mm")'
                f'=TEXT($A{row},"yyyy-mm"))*${helper_income_col}${helper_start}:${helper_income_col}${helper_end})'
            ),
        )
        ws.cell(row=row, column=4, value=f"=C{row}-B{row}")
        for col_letter in ("B", "C", "D"):
            ws[f"{col_letter}{row}"].number_format = base.MONEY_FORMAT

    sum_row = end + 1
    base.add_sum_row(ws, sum_row, 1, ["B", "C", "D"], header_row, row_count)

    base.set_widths(ws, [12, 14, 16, 14])
    ws.column_dimensions["F"].width = 20
    ws.freeze_panes = f"A{header_row + 1}"


def build_거래처_view(wb: Workbook, sheet_name: str, sort_col_formula: str, description: str) -> None:
    ws = wb.create_sheet(sheet_name)
    last_col = len(거래처_HEADERS)
    header_row = 4

    base.write_title_block(
        ws, sheet_name, description,
        "※ 원본 거래처 관리 파일의 '거래처 관리 대장' 탭을 IMPORTRANGE로 실시간 정렬해서 보여줍니다. 최초 1회 '액세스 허용' 클릭이 필요합니다.",
        last_col,
    )
    base.write_header(ws, header_row, 거래처_HEADERS)
    ws.cell(
        row=header_row + 1, column=1,
        value=(
            f'=SORT(IMPORTRANGE("{MASTER_SHEET_URL}","{MASTER_VENDOR_DATA_RANGE}"), {sort_col_formula}, TRUE)'
        ),
    )
    base.set_widths(ws, [6, 24, 20, 20, 18, 10, 26, 22, 14, 12, 12, 24])
    ws.freeze_panes = f"A{header_row + 1}"


def main() -> None:
    wb = Workbook()
    default_sheet = wb.active
    build_거래처_view(
        wb, "거래처 관리 (위치순)", "3",
        "위치정보(주소) 가나다순 자동 정렬 뷰",
    )
    build_거래처_view(
        wb, "거래처 관리 (발행일순)",
        'IFERROR(VALUE(REGEXREPLACE(IMPORTRANGE("' + MASTER_SHEET_URL + '","' + MASTER_VENDOR_TAB + '!J5:J34"),"[^0-9]","")),9999)',
        "세금계산서 발행일(N일) 오름차순 자동 정렬 뷰",
    )
    base.build_지출관리(wb)
    build_월별순수익금(wb)
    wb.remove(default_sheet)

    wb.save(OUTPUT_PATH)
    print(f"생성 완료: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
